"""Генератор електронної таблиці CRC-карток docs/crc.xlsx.

Створює стилізовану книгу Excel з 8 сутностями предметної області
відповідно до вимог FR-017, FR-018 та contracts/crc-contract.md.
"""

import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_FILE = PROJECT_ROOT / "docs" / "crc.xlsx"

CRC_CARDS_DATA = [
    {
        "class_name": "IDocumentPackageHandler",
        "type": "Interface",
        "role": "Контракт обробника товаросупровідних документів митного оформлення.",
        "responsibilities": [
            "1. Визначати уніфікований контракт перевірки та прийняття товаросупровідних документів (processDocumentPackage).",
            "2. Забезпечувати поліморфну обробку пакетів документів для різних ролей учасників.",
        ],
        "collaborators": [
            "DocumentPackage",
        ],
    },
    {
        "class_name": "CustomsParticipant",
        "type": "Abstract Class",
        "role": "Базовий суб'єкт митних правовідносин та взаємодій.",
        "responsibilities": [
            "1. Зберігати базові атрибути суб'єкта (ідентифікатор, найменування, email).",
            "2. Надавати гарантовану незмінну ідентифікацію учасника (final getParticipantInfo).",
            "3. Декларувати обов'язок реалізації обробки документів (processDocumentPackage).",
        ],
        "collaborators": [
            "IDocumentPackageHandler",
            "DocumentPackage",
        ],
    },
    {
        "class_name": "Importer",
        "type": "Concrete Class",
        "role": "Суб'єкт зовнішньоекономічної діяльності — власник або одержувач вантажу.",
        "responsibilities": [
            "1. Зберігати реквізити підприємства-імпортера (код ЄДРПОУ, податкова адреса).",
            "2. Формувати первинний пакет товаросупровідних документів (processDocumentPackage).",
            "3. Ініціювати передачу пакета документів митному брокеру (submitDocuments).",
        ],
        "collaborators": [
            "CustomsParticipant",
            "CustomsBroker",
            "DocumentPackage",
        ],
    },
    {
        "class_name": "CustomsBroker",
        "type": "Concrete Class",
        "role": "Ліцензований представник, що здійснює декларування товарів від імені клієнта.",
        "responsibilities": [
            "1. Зберігати дані ліцензії брокера та комерційні умови надання послуг.",
            "2. Проводити експертизу комплектності товаросупровідних документів (processDocumentPackage).",
            "3. Створювати проект електронної митної декларації (createDeclaration).",
            "4. Здійснювати класифікацію товарних позицій за класифікатором (classifyItem).",
            "5. Подавати завірену декларацію посадовій особі митниці (submitDeclaration).",
        ],
        "collaborators": [
            "CustomsParticipant",
            "Importer",
            "CustomsDeclaration",
            "CommodityCodeRegistry",
            "CustomsInspector",
            "DocumentPackage",
        ],
    },
    {
        "class_name": "CustomsInspector",
        "type": "Concrete Class",
        "role": "Посадова особа митного органу, уповноважена здійснювати контроль та оформлення.",
        "responsibilities": [
            "1. Зберігати особисті службові реквізити (номер печатки/жетона, код поста).",
            "2. Здійснювати нормативну верифікацію документів (processDocumentPackage).",
            "3. Призначати форму митного контролю за профілем ризику (assignInspection).",
            "4. Проводити перевірку декларації та розрахунку платежів (verifyDocuments).",
            "5. Приймати остаточне юридичне рішення про випуск товарів (releaseCargo).",
        ],
        "collaborators": [
            "CustomsParticipant",
            "CustomsDeclaration",
            "DocumentPackage",
        ],
    },
    {
        "class_name": "CustomsDeclaration",
        "type": "Concrete Class",
        "role": "Електронний документ митної декларації, що містить відомості про вантаж, вартість та платежі.",
        "responsibilities": [
            "1. Агрегувати відомості про учасників, товарні позиції та статус процедури.",
            "2. Керувати складом товарних позицій (композиція addItem, getItems).",
            "3. Виконувати розрахунок суми митних платежів за базовою ставкою (calculateDuties(rate)).",
            "4. Виконувати розрахунок платежів із врахуванням преференційних знижок (calculateDuties(rate, discount)).",
            "5. Верифікувати коректність кодів через зовнішній реєстр (validateCommodityCodes).",
            "6. Фіксувати сплату платежів та оновлення статусів життєвого циклу (markDutiesPaid, updateStatus).",
        ],
        "collaborators": [
            "DeclarationItem",
            "CommodityCodeRegistry",
        ],
    },
    {
        "class_name": "DeclarationItem",
        "type": "Concrete Class",
        "role": "Окрема товарна позиція партії вантажу, що декларується.",
        "responsibilities": [
            "1. Зберігати комерційний опис, код УКТЗЕД, фактурну вартість, вагу та країну походження.",
            "2. Надавати форматовані дані для нарахування податків (getItemDetails).",
            "3. Гарантувати цілісність даних у межах батьківської декларації.",
        ],
        "collaborators": [
            "CustomsDeclaration",
        ],
    },
    {
        "class_name": "CommodityCodeRegistry",
        "type": "Concrete Class",
        "role": "Довідник Української класифікації товарів ЗЕД та нормативних обмежень.",
        "responsibilities": [
            "1. Зберігати актуальні ставки ввізного мита для кодів класифікатора.",
            "2. Перевіряти коректність синтаксичного формату коду товару (validateCodeFormat).",
            "3. Надавати ставки мита за запитом (getTariffRate).",
            "4. Визначати перелік обов'язкових дозвільних документів та нетарифних обмежень (getRequiredPermits).",
        ],
        "collaborators": [
            "CustomsDeclaration",
            "CustomsBroker",
        ],
    },
]


def generate_crc():
    print("=== Генерація книги CRC-карток docs/crc.xlsx ===")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CRC Cards"
    ws.views.sheetView[0].showGridLines = True

    # Стилі
    font_header = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    font_subheader = Font(name="Segoe UI", size=11, bold=True, color="000000")
    font_data = Font(name="Segoe UI", size=10, color="000000")
    font_role = Font(name="Segoe UI", size=9, italic=True, color="E0E0E0")

    fill_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_subheader = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_even = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="8EA9DB")
    thin_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

    current_row = 2

    # Заголовок документа
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
    doc_title_cell = ws.cell(row=current_row, column=1)
    doc_title_cell.value = "КАТАЛОГ CRC-КАРТОК СИСТЕМИ МИТНОГО БРОКЕРА (CustomsBroker)"
    doc_title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    doc_title_cell.fill = PatternFill(start_color="0D233A", end_color="0D233A", fill_type="solid")
    doc_title_cell.alignment = align_center
    ws.row_dimensions[current_row].height = 32
    ws.cell(row=current_row, column=2).border = thin_border
    doc_title_cell.border = thin_border
    current_row += 2

    for card in CRC_CARDS_DATA:
        card_start_row = current_row

        # 1. Заголовок картки класу
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
        header_cell = ws.cell(row=current_row, column=1)
        header_cell.value = f"Клас: {card['class_name']} ({card['type']})\nРоль: {card['role']}"
        header_cell.font = font_header
        header_cell.fill = fill_header
        header_cell.alignment = align_left
        ws.row_dimensions[current_row].height = 42

        # Оформлення правої клітинки мержу рамкою
        ws.cell(row=current_row, column=2).border = thin_border
        header_cell.border = thin_border
        current_row += 1

        # 2. Підзаголовки колонок
        col1_head = ws.cell(row=current_row, column=1)
        col1_head.value = "Обов'язки (Responsibilities)"
        col1_head.font = font_subheader
        col1_head.fill = fill_subheader
        col1_head.alignment = align_center
        col1_head.border = thin_border

        col2_head = ws.cell(row=current_row, column=2)
        col2_head.value = "Співпраця (Collaborators)"
        col2_head.font = font_subheader
        col2_head.fill = fill_subheader
        col2_head.alignment = align_center
        col2_head.border = thin_border
        ws.row_dimensions[current_row].height = 24
        current_row += 1

        # 3. Рядки даних
        resps = card["responsibilities"]
        collabs = card["collaborators"]
        num_rows = max(len(resps), len(collabs))

        for i in range(num_rows):
            r_val = resps[i] if i < len(resps) else ""
            c_val = collabs[i] if i < len(collabs) else ""

            cell_r = ws.cell(row=current_row, column=1, value=r_val)
            cell_r.font = font_data
            cell_r.alignment = align_left
            cell_r.border = thin_border
            cell_r.fill = fill_even if (i % 2 == 1) else fill_white

            cell_c = ws.cell(row=current_row, column=2, value=c_val)
            cell_c.font = font_data
            cell_c.alignment = align_left
            cell_c.border = thin_border
            cell_c.fill = fill_even if (i % 2 == 1) else fill_white

            ws.row_dimensions[current_row].height = 28 if len(r_val) > 40 else 22
            current_row += 1

        # Відступ між картками
        current_row += 1

    # Встановлення ширини стовпчиків
    ws.column_dimensions["A"].width = 68
    ws.column_dimensions["B"].width = 42

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    print(f"[SUCCESS] Згенеровано {OUTPUT_FILE} з {len(CRC_CARDS_DATA)} CRC-картками.")


if __name__ == "__main__":
    generate_crc()
